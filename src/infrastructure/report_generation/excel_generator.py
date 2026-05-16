"""
Excel report generator.

Creates formatted Excel spreadsheets for tabular data.
"""

from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Dict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.config.logging_config import get_logger
from src.domain.interfaces.report_generator import ReportGenerator
from src.domain.models.response import Response
from src.utils.exceptions import ReportGenerationError

logger = get_logger(__name__)


class ExcelReportGenerator(ReportGenerator):
    """
    Generates formatted Excel reports using openpyxl.
    """

    def __init__(self):
        """Initialize Excel generator."""
        logger.info("excel_generator_initialized")

    def generate_excel_report(
        self,
        data: Dict[str, Any],
        query_text: str,
        metadata: Dict[str, Any],
    ) -> bytes:
        """
        Generate an Excel report from structured data.

        Args:
            data: Structured data to format as Excel
            query_text: Original query
            metadata: Additional report metadata

        Returns:
            bytes: Excel file content

        Raises:
            ReportGenerationError: If generation fails
        """
        logger.info("generating_excel_report", query=query_text[:50])

        try:
            # Create workbook
            wb = Workbook()

            # Remove default sheet
            wb.remove(wb.active)

            # Sheet 1: Summary
            ws_summary = wb.create_sheet("Summary", 0)
            self._build_summary_sheet(ws_summary, query_text, metadata)

            # Sheet 2: Data
            if data:
                ws_data = wb.create_sheet("Data", 1)
                self._build_data_sheet(ws_data, data)

            # Sheet 3: Metadata
            ws_meta = wb.create_sheet("Metadata", 2)
            self._build_metadata_sheet(ws_meta, metadata)

            # Save to bytes
            buffer = BytesIO()
            wb.save(buffer)
            excel_bytes = buffer.getvalue()

            logger.info("excel_report_generated", size_bytes=len(excel_bytes))

            return excel_bytes

        except Exception as e:
            logger.error("excel_generation_failed", error=str(e))
            raise ReportGenerationError(f"Failed to generate Excel: {e}")

    def _build_summary_sheet(
        self,
        ws,
        query_text: str,
        metadata: Dict[str, Any],
    ):
        """Build summary sheet."""
        # Title
        ws.append(["Financial Analysis Report"])
        ws["A1"].font = Font(size=16, bold=True)

        # Timestamp
        ws.append([f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
        ws.append([])

        # Query
        ws.append(["Query:"])
        ws["A4"].font = Font(bold=True)
        ws.append([query_text])
        ws.merge_cells("A5:D5")
        ws.append([])

        # Apply styles
        self._apply_header_style(ws, "A1")

    def _build_data_sheet(self, ws, data: Dict[str, Any]):
        """Build data sheet with structured data."""
        # Handle different data structures
        if "rows" in data:
            self._build_table_from_rows(ws, data["rows"])
        elif "metrics" in data:
            self._build_table_from_metrics(ws, data["metrics"])
        else:
            # Generic key-value display
            self._build_key_value_table(ws, data)

    def _build_table_from_rows(self, ws, rows: list):
        """Build table from row data."""
        if not rows:
            ws.append(["No data available"])
            return

        # Get headers from first row
        if isinstance(rows[0], dict):
            headers = list(rows[0].keys())
            ws.append(headers)

            # Add data rows
            for row in rows:
                ws.append([row.get(h, "") for h in headers])

        else:
            # Simple list
            for row in rows:
                ws.append([row] if not isinstance(row, (list, tuple)) else list(row))

        # Style header row
        self._style_header_row(ws, 1, len(headers) if rows else 1)
        self._auto_size_columns(ws)

    def _build_table_from_metrics(self, ws, metrics: Dict):
        """Build table from metrics dictionary."""
        # Header
        ws.append(["Metric", "Value"])

        # Data
        for key, value in metrics.items():
            ws.append([key, str(value)])

        # Style
        self._style_header_row(ws, 1, 2)
        self._auto_size_columns(ws)

    def _build_key_value_table(self, ws, data: Dict):
        """Build generic key-value table."""
        ws.append(["Key", "Value"])

        for key, value in data.items():
            # Handle nested structures
            if isinstance(value, (dict, list)):
                value = str(value)

            ws.append([key, value])

        self._style_header_row(ws, 1, 2)
        self._auto_size_columns(ws)

    def _build_metadata_sheet(self, ws, metadata: Dict[str, Any]):
        """Build metadata sheet."""
        ws.append(["Report Metadata"])
        ws["A1"].font = Font(size=14, bold=True)
        ws.append([])

        # Add metadata
        for key, value in metadata.items():
            ws.append([key, str(value)])

        self._auto_size_columns(ws)

    def _style_header_row(self, ws, row_num: int, col_count: int):
        """Apply styling to header row."""
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def _apply_header_style(self, ws, cell_ref: str):
        """Apply header style to a cell."""
        cell = ws[cell_ref]
        cell.font = Font(size=16, bold=True, color="1F4E78")

    def _auto_size_columns(self, ws):
        """Auto-size columns based on content."""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)

            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass

            adjusted_width = min(max_length + 2, 50)  # Cap at 50
            ws.column_dimensions[column_letter].width = adjusted_width

    def generate_pdf_report(
        self,
        response: Response,
        query_text: str,
        metadata: Dict[str, Any],
    ) -> bytes:
        """Not implemented in Excel generator."""
        raise NotImplementedError("Use PDFReportGenerator for PDF reports")

    def save_report(
        self,
        content: bytes,
        filename: str,
        output_dir: str,
    ) -> str:
        """
        Save report to file system.

        Args:
            content: Report file content
            filename: Name for the file
            output_dir: Directory to save in

        Returns:
            str: Path to saved file
        """
        try:
            output_path = Path(output_dir)
            output_path.mkdir(parents=True, exist_ok=True)

            file_path = output_path / filename

            with open(file_path, "wb") as f:
                f.write(content)

            logger.info("report_saved", file_path=str(file_path))

            return str(file_path)

        except Exception as e:
            logger.error("report_save_failed", error=str(e))
            raise ReportGenerationError(f"Failed to save report: {e}")
